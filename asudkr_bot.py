#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Telegram-бот для АСУ ДКР (asudkros.railways.kz).

Как работает:
  1. Вы пишете боту дату или период (например 31.08.2026, "сегодня", "вчера",
     или 01.08.2026-31.08.2026).
  2. Бот заходит в АСУ ДКР под вашим логином/паролем и присылает вам КАПЧУ (картинку).
  3. Вы отвечаете кодом с картинки.
  4. Бот забирает все накладные за дату/период (номер, станции, отправитель/получатель,
     груз, вес, вагоны) и присылает готовый Excel.

Запуск:
  pip install pyTelegramBotAPI requests openpyxl
  Задайте настройки ниже (или через переменные окружения) и:
  python asudkr_bot.py
"""

import os
import io
import re
import sys
import base64
from datetime import datetime, date, timedelta

import requests
import telebot
from telebot import types
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ======================= НАСТРОЙКИ =======================
# Настройки читаются из файла config.txt в той же папке (правьте его в Блокноте).
# Если файла нет — берутся из переменных окружения (для сервера).
def _load_config():
    cfg = {}
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.txt")
    if os.path.exists(path):
        for line in open(path, encoding="utf-8-sig"):
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            cfg[k.strip().upper()] = v.strip()
    return cfg

_CFG = _load_config()
def _get(key, default=""):
    return _CFG.get(key) or os.environ.get(key, default)

BOT_TOKEN       = _get("BOT_TOKEN",       "ВСТАВЬТЕ_ТОКЕН_ОТ_BOTFATHER")
ASUDKR_USER     = _get("ASUDKR_USER",     "ВАШ_ЛОГИН_АСУ_ДКР")
ASUDKR_PASSWORD = _get("ASUDKR_PASSWORD", "ВАШ_ПАРОЛЬ_АСУ_ДКР")
ALLOWED_IDS     = _get("ALLOWED_IDS",     "")   # пусто = разрешить всем (бот подскажет ваш ID)
# =========================================================

BASE = "https://asudkros.railways.kz"
bot = telebot.TeleBot(BOT_TOKEN)


def log(msg):
    print(msg, flush=True)

# состояние диалога по каждому чату
STATE = {}   # chat_id -> {"step":..., "period":(date,date), "session":requests.Session(), "login_ctx":{...}}


def allowed(uid):
    if not ALLOWED_IDS.strip():
        return True
    ids = [x.strip() for x in ALLOWED_IDS.replace(";", ",").split(",") if x.strip()]
    return str(uid) in ids


# ------------------------- РАЗБОР ДАТЫ / ПЕРИОДА -------------------------
def parse_user_date(text):
    t = text.strip().lower()
    today = date.today()
    if t in ("сегодня", "бүгін", "today"):
        return today
    if t in ("вчера", "кеше", "yesterday"):
        return today - timedelta(days=1)
    t = t.replace("/", ".").replace("-", ".")
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%d.%m"):
        try:
            d = datetime.strptime(t, fmt).date()
            if fmt == "%d.%m":
                d = d.replace(year=today.year)
            return d
        except ValueError:
            continue
    return None


def parse_user_period(text):
    """Принимает одну дату или период: "01.08.2026-31.08.2026",
    "01.08.2026 - 31.08.2026", "01.08.2026 по 31.08.2026".
    Возвращает (start, end) или None."""
    t = text.strip()

    def _try_split(left, right):
        left, right = left.strip(), right.strip()
        if not left or not right:
            return None
        d1, d2 = parse_user_date(left), parse_user_date(right)
        if d1 and d2:
            return (d1, d2) if d1 <= d2 else (d2, d1)
        return None

    for sep in (" по ", " - ", "–", "—"):
        if sep in t:
            res = _try_split(*t.split(sep, 1))
            if res:
                return res

    # голый дефис-разделитель — только если по обе стороны похоже на дату с точками
    # (иначе "31-08-2026" читалось бы как период, а не как дата с дефисами)
    if "-" in t:
        left, right = t.split("-", 1)
        if "." in left and "." in right:
            res = _try_split(left, right)
            if res:
                return res

    d = parse_user_date(t)
    if d:
        return (d, d)
    return None


_MONTHS = {m: i for i, m in enumerate(
    ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"], 1)}

def invoice_date_to_date(s):
    # "Aug 31, 2026, 1:06:14 PM" -> date(2026,8,31)
    m = re.match(r"([A-Za-z]{3})\s+(\d{1,2}),\s*(\d{4})", s or "")
    if not m:
        return None
    mon = _MONTHS.get(m.group(1))
    if not mon:
        return None
    return date(int(m.group(3)), mon, int(m.group(2)))


# ------------------------- ПАРСЕР .txt (EDIFACT/СМГС) -------------------------
def parse_edi(text):
    segs = re.findall(r"[A-Z]{3}\+[^']*'", text or "")
    weight, wagons = "", []
    hs_code, etsng_code, cargo_name = "", "", ""
    for s in segs:
        p = s.rstrip("'").split("+")
        tag = p[0]
        if tag == "MEA" and len(p) > 3 and p[1] == "WT" and p[2] == "G" and "KGM" in p[3] and not weight:
            weight = p[3].split(":")[-1]
        elif tag == "EQD" and len(p) > 2:
            num = p[2].split(":")[0]
            if num and num not in wagons:
                wagons.append(num)
        elif tag == "PIA" and len(p) > 2:
            # PIA+5+<код>:<квалификатор>::<...>  — HS = код ТН ВЭД, ET = код ЕТСНГ/ГНГ
            fields = p[2].split(":")
            code, qual = fields[0], (fields[1] if len(fields) > 1 else "")
            if qual == "HS" and not hs_code:
                hs_code = code
            elif qual == "ET" and not etsng_code:
                etsng_code = code
        elif tag == "FTX" and len(p) > 1 and p[1] == "AAA" and not cargo_name:
            # FTX+AAA+++<наименование груза>
            cargo_name = p[-1]
    return weight, wagons, hs_code, etsng_code, cargo_name


def format_goods_detail(hs_code, etsng_code, cargo_name):
    code_part = hs_code
    if etsng_code:
        code_part = "%s (%s)" % (hs_code, etsng_code) if hs_code else "(%s)" % etsng_code
    if code_part and cargo_name:
        return "%s, %s" % (code_part, cargo_name)
    return code_part or cargo_name


# ------------------------- РАБОТА С АСУ ДКР -------------------------
def api_get_captcha(sess):
    r = sess.post(BASE + "/api/sso/auth/getcaptcha", timeout=30)
    r.raise_for_status()
    raw = r.text.strip().strip('"')
    if raw.startswith("data:image"):
        raw = raw.split(",", 1)[1]
    try:
        return base64.b64decode(raw)
    except Exception:
        return r.content  # вдруг вернулись сырые байты картинки


def api_login(sess, captcha_text):
    body = {"username": ASUDKR_USER, "password": ASUDKR_PASSWORD, "captchaText": captcha_text}
    r = sess.post(BASE + "/api/sso/auth/OUT", json=body, timeout=30)
    try:
        j = r.json()
    except Exception:
        return None, "Сервер вернул неожиданный ответ при входе (код %s)." % r.status_code
    if isinstance(j, dict) and j.get("token"):
        return j, None
    if isinstance(j, dict) and j.get("msg") == "2faSent":
        return {"_2fa": True, "ctx": j}, None
    msg = (j.get("msg") or j.get("error") or "неверный логин/пароль или код") if isinstance(j, dict) else "ошибка"
    return None, "Не удалось войти: %s" % msg


def api_valid_2fa(sess, ctx, code):
    body = {"username": ASUDKR_USER, "password": ASUDKR_PASSWORD, "code": str(code).strip()}
    if isinstance(ctx, dict):
        body.update({k: v for k, v in ctx.items() if k not in ("msg",)})
    r = sess.post(BASE + "/api/sso/auth/validauthcode", json=body, timeout=30)
    try:
        j = r.json()
    except Exception:
        return None, "Неожиданный ответ при подтверждении кода."
    if j.get("token"):
        return j, None
    return None, "Код подтверждения не принят."


def auth_headers(identity):
    return {"Authorization": (identity.get("tokenType") or "Bearer") + " " + identity["token"]}


def _epd_invoice_request(sess, H, page, limit, start, end):
    # dateToStart/dateToStartTo — реальные параметры фильтра "Дата принятия к
    # перевозке" (эндпоинт "epd"/"getInvoice"), формат dd.MM.yyyy — взято из
    # исходного JS-кода сайта (Angular filterForm), т.к. документации к API нет.
    url = (BASE + "/api/epd/searchdata?method=getInvoice"
           "&invoice=all&railDivUn=&page=%d&limit=%d&dateToStart=%s&dateToStartTo=%s"
           % (page, limit, start.strftime("%d.%m.%Y"), end.strftime("%d.%m.%Y")))
    r = sess.get(url, headers=H, timeout=40)
    if r.status_code != 200:
        return None
    return r.json() or {}


def fetch_invoices_for_period(sess, identity, start, end):
    H = auth_headers(identity)
    rows, seen_ids = [], set()
    total = None
    page, limit = 1, 100
    PAGE_LIMIT = 200  # защита от зацикливания

    while page <= PAGE_LIMIT:
        data = _epd_invoice_request(sess, H, page, limit, start, end)
        if not data:
            break
        chunk = data.get("rows", []) or []
        if page == 1:
            for key in ("total", "totalCount", "count", "recordsTotal", "totalRows"):
                if isinstance(data.get(key), int):
                    total = data[key]
                    break
            log("[fetch] dateToStart=%s dateToStartTo=%s total=%s chunk=%d"
                % (start.strftime("%d.%m.%Y"), end.strftime("%d.%m.%Y"), total, len(chunk)))
        if not chunk:
            break
        new_ids = [x.get("id") for x in chunk]
        if new_ids and all(i in seen_ids for i in new_ids):
            log("[fetch] page %d: сервер вернул уже виденные записи, стоп" % page)
            break
        seen_ids.update(new_ids)
        rows.extend(chunk)
        if total is not None and len(rows) >= total:
            break
        if len(chunk) < limit:
            break
        page += 1
    log("[fetch] итого_получено=%d (сервер сообщал total=%s)" % (len(rows), total))

    # на всякий случай ещё раз фильтруем по дате на своей стороне
    # (сервер уже должен был отфильтровать по dateToStart/dateToStartTo)
    keep = []
    unparsed = 0
    for x in rows:
        raw = x.get("createDate")
        d = invoice_date_to_date(raw)
        if d is None:
            unparsed += 1
            continue
        if start <= d <= end:
            keep.append((d, x))
    keep.sort(key=lambda p: p[0])
    log("[fetch] собрано_строк=%d не_распознана_дата=%d подходит_под_период=%d период=%s..%s пример_дат=%s"
        % (len(rows), unparsed, len(keep), start, end,
           [x.get("createDate") for x in rows[:3]]))
    result = []
    for d, x in keep:
        weight, wagons = "", []
        hs_code, etsng_code, cargo_name = "", "", ""
        try:
            fr = sess.get(BASE + "/api/epd/searchdata?method=getMessageFile&invoiceId=%s" % x["id"],
                          headers=H, timeout=40)
            if fr.status_code == 200:
                weight, wagons, hs_code, etsng_code, cargo_name = parse_edi(fr.text)
        except Exception:
            pass
        if not wagons:
            fv, lv = x.get("firstVag"), x.get("lastVag")
            wagons = [v for v in {fv, lv} if v]
        result.append({
            "date": d, "num": x.get("invoiceNum", ""), "staSend": x.get("staSend", ""),
            "staDest": x.get("stationDest", ""), "sender": x.get("sender", ""),
            "receiver": re.sub(r"\s+", " ", x.get("receiver", "") or ""),
            "gruz": x.get("gruzName", ""), "weight": weight, "wagons": wagons,
            "goods_detail": format_goods_detail(hs_code, etsng_code, cargo_name),
        })
    return result


# ------------------------- ПОСТРОЕНИЕ EXCEL -------------------------
def build_excel(records, start, end):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Накладные"
    headers = ["№", "Дата", "Номер накладной", "Станция отправитель", "Станция получатель",
               "Грузоотправитель", "Грузополучатель", "Груз", "Код ТН ВЭД (ЕТСНГ), наименование",
               "Вес, кг", "Кол-во вагонов", "Номера вагонов"]
    ws.append(headers)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    total_w = 0
    for i, r in enumerate(records, 1):
        w = int(r["weight"]) if str(r["weight"]).isdigit() else 0
        total_w += w
        ws.append([i, r["date"].strftime("%d.%m.%Y"), r["num"], r["staSend"], r["staDest"],
                   r["sender"], r["receiver"], r["gruz"], r["goods_detail"], w,
                   len(r["wagons"]), ", ".join(r["wagons"])])
    for row in ws.iter_rows(min_row=2, max_row=1 + len(records)):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True)
        row[9].number_format = "#,##0"
    if records:
        tr = ws.max_row + 1
        ws.cell(tr, 9, "ИТОГО:").font = Font(name="Arial", bold=True)
        ws.cell(tr, 9).alignment = Alignment(horizontal="right")
        c10 = ws.cell(tr, 10, total_w); c10.font = Font(name="Arial", bold=True); c10.number_format = "#,##0"
        ws.cell(tr, 11, sum(len(r["wagons"]) for r in records)).font = Font(name="Arial", bold=True)
    for i, wdt in enumerate([4, 12, 16, 18, 20, 24, 32, 10, 40, 12, 13, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    if start == end:
        buf.name = "Накладные_%s.xlsx" % start.strftime("%Y%m%d")
    else:
        buf.name = "Накладные_%s-%s.xlsx" % (start.strftime("%Y%m%d"), end.strftime("%Y%m%d"))
    return buf, total_w


# ------------------------- ХЭНДЛЕРЫ TELEGRAM -------------------------
@bot.message_handler(commands=["start", "help"])
def cmd_start(m):
    uid = m.from_user.id
    if not allowed(uid):
        return bot.reply_to(m, "Извините, этот бот приватный.\nВаш ID: %s" % uid)
    hint = ""
    if not ALLOWED_IDS.strip():
        hint = ("\n\nℹ️ Ваш Telegram ID: %s\n"
                "Впишите его в config.txt (строка ALLOWED_IDS=...), "
                "чтобы ботом могли пользоваться только вы." % uid)
    bot.reply_to(m,
        "Здравствуйте! Я собираю накладные из АСУ ДКР в Excel.\n\n"
        "Напишите мне дату или период, например:\n"
        "• 31.08.2026\n• сегодня\n• вчера\n• 01.08.2026-31.08.2026\n\n"
        "Я зайду в систему, пришлю вам код с картинки для подтверждения, "
        "и верну готовую таблицу за эту дату/период." + hint)


def _fmt_period(period):
    start, end = period
    if start == end:
        return start.strftime("%d.%m.%Y")
    return "%s — %s" % (start.strftime("%d.%m.%Y"), end.strftime("%d.%m.%Y"))


def start_login(chat_id, period):
    sess = requests.Session()
    try:
        img = api_get_captcha(sess)
    except Exception as e:
        bot.send_message(chat_id, "Не удалось получить капчу: %s" % e)
        return
    STATE[chat_id] = {"step": "captcha", "period": period, "session": sess}
    bot.send_photo(chat_id, img,
                   caption="Период: %s\nВведите код с картинки:" % _fmt_period(period))


def do_export(chat_id):
    st = STATE.get(chat_id, {})
    sess, identity, period = st.get("session"), st.get("identity"), st.get("period")
    start, end = period
    bot.send_message(chat_id, "Вошёл. Собираю накладные за %s…" % _fmt_period(period))
    try:
        records = fetch_invoices_for_period(sess, identity, start, end)
    except Exception as e:
        bot.send_message(chat_id, "Ошибка при получении данных: %s" % e)
        STATE.pop(chat_id, None); return
    if not records:
        bot.send_message(chat_id, "За %s накладных не найдено." % _fmt_period(period))
        STATE.pop(chat_id, None); return
    buf, total_w = build_excel(records, start, end)
    bot.send_document(chat_id, buf,
        caption="Готово: %d накладных, вес %s кг." % (len(records), format(total_w, ",d").replace(",", " ")))
    STATE.pop(chat_id, None)


@bot.message_handler(func=lambda m: True)
def on_text(m):
    if not allowed(m.from_user.id):
        return bot.reply_to(m, "Извините, этот бот приватный.")
    chat_id = m.chat.id
    st = STATE.get(chat_id)
    text = (m.text or "").strip()

    # ожидаем код капчи
    if st and st.get("step") == "captcha":
        identity, err = api_login(st["session"], text)
        if err:
            bot.reply_to(m, err + "\nПопробуйте снова — пришлите дату ещё раз.")
            STATE.pop(chat_id, None); return
        if identity.get("_2fa"):
            st["step"] = "2fa"; st["ctx"] = identity["ctx"]
            return bot.reply_to(m, "Отправлен код подтверждения. Введите его:")
        st["identity"] = identity; st["step"] = "ready"
        return do_export(chat_id)

    # ожидаем код 2FA
    if st and st.get("step") == "2fa":
        identity, err = api_valid_2fa(st["session"], st.get("ctx"), text)
        if err:
            bot.reply_to(m, err + "\nПопробуйте снова — пришлите дату ещё раз.")
            STATE.pop(chat_id, None); return
        st["identity"] = identity; st["step"] = "ready"
        return do_export(chat_id)

    # иначе это должна быть дата или период
    period = parse_user_period(text)
    if not period:
        return bot.reply_to(m, "Не поняла дату/период. Пример: 31.08.2026, сегодня, вчера, "
                                "01.08.2026-31.08.2026.")
    start_login(chat_id, period)


if __name__ == "__main__":
    print("АСУ ДКР бот запущен. Ожидаю сообщения…")
    bot.infinity_polling(skip_pending=True)
